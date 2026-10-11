import base64
import csv
import io
import os
import secrets
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

import fitz
from openpyxl import load_workbook
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel

from .pdf_processor import analyze_pdf, preview_pdf, process_pdf

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)
MAX_FILE_SIZE = 25 * 1024 * 1024
MAX_CSV_SIZE = 5 * 1024 * 1024
SESSION_TTL = 60 * 60

app = FastAPI(title="Amazon SKU Label Printer", version="1.0.0")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "app" / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "app" / "templates"))

# Short-lived in-memory state avoids persistent storage of customer label PDFs.
sessions: Dict[str, Dict[str, Any]] = {}


class GenerateRequest(BaseModel):
    file_id: str
    skus: List[Optional[str]]
    labels_only: bool = True
    sku_x: float = 300.0
    sku_y: Optional[float] = None
    sku_font_size: float = 10.0


def cleanup_sessions() -> None:
    cutoff = time.time() - SESSION_TTL
    expired = [key for key, value in sessions.items() if value.get("created", 0) < cutoff]
    for key in expired:
        sessions.pop(key, None)


def validate_pdf(data: bytes, filename: str) -> None:
    if not filename.lower().endswith(".pdf"):
        raise HTTPException(400, "Please upload a valid PDF file.")
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(413, "PDF is too large. Maximum supported size is 25 MB.")
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        if doc.is_encrypted and not doc.authenticate(""):
            doc.close()
            raise HTTPException(400, "Password-protected PDFs are not supported.")
        if doc.page_count == 0:
            doc.close()
            raise HTTPException(400, "The PDF contains no pages.")
        doc.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "The uploaded PDF is corrupted or unsupported.") from exc


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    cleanup_sessions()
    return templates.TemplateResponse("index.html", {"request": request})


@app.get("/health")
async def health():
    return {"status": "ok"}


@app.post("/upload")
async def upload(file: UploadFile = File(...)):
    cleanup_sessions()
    data = await file.read()
    validate_pdf(data, file.filename or "upload.pdf")
    file_id = secrets.token_urlsafe(16)
    doc = fitz.open(stream=data, filetype="pdf")
    pages = doc.page_count
    doc.close()
    sessions[file_id] = {
        "created": time.time(),
        "original": data,
        "filename": file.filename or "amazon_labels.pdf",
        "mapping": {},
        "mapping_mode": "old_new",
    }
    return {"file_id": file_id, "filename": file.filename, "size": len(data), "pages": pages}


@app.post("/upload-csv")
async def upload_csv(file_id: str, file: UploadFile = File(...)):
    cleanup_sessions()
    state = sessions.get(file_id)
    if not state:
        raise HTTPException(404, "Upload session not found. Please upload the PDF again.")
    name = (file.filename or "").lower()
    if not (name.endswith(".csv") or name.endswith(".xlsx")):
        raise HTTPException(400, "Please upload a CSV or Excel (.xlsx) file.")
    data = await file.read()
    if len(data) > MAX_CSV_SIZE:
        raise HTTPException(413, "Mapping file is too large. Maximum supported size is 5 MB.")

    rows = []
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            workbook.close()
            if values:
                headers = [str(v or "").strip().lower().replace("\ufeff", "") for v in values[0]]
                rows = [dict(zip(headers, row)) for row in values[1:]]
        except Exception as exc:
            raise HTTPException(400, "Could not read Excel file. Please upload a valid .xlsx file.") from exc
    else:
        encodings = ["utf-8-sig", "utf-8", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "latin-1"]
        parsed = None
        last_error = None
        for encoding in encodings:
            try:
                text = data.decode(encoding)
                try:
                    dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|:")
                except csv.Error:
                    dialect = csv.excel
                candidate = list(csv.DictReader(io.StringIO(text), dialect=dialect))
                if candidate:
                    headers = {str(k).strip().lower().replace("\ufeff", ""): k for k in candidate[0].keys() if k}
                    if (("old sku" in headers and "new sku" in headers) or
                        ("product_name" in headers and "sku" in headers) or
                        ("product name" in headers and "sku" in headers)):
                        parsed = candidate
                        break
            except (UnicodeDecodeError, UnicodeError, csv.Error, ValueError) as exc:
                last_error = exc
        if parsed is None:
            raise HTTPException(400, "Use columns Old SKU and New SKU (or legacy Product Name and SKU).") from last_error
        rows = [{str(k or "").strip().lower().replace("\ufeff", ""): v for k, v in row.items()} for row in parsed]

    if not rows:
        raise HTTPException(400, "The mapping file has no data rows.")
    headers = set(rows[0].keys())
    old_key = next((k for k in ("old sku", "old_sku", "oldsku") if k in headers), None)
    new_key = next((k for k in ("new sku", "new_sku", "newsku") if k in headers), None)
    if old_key and new_key:
        mapping_mode = "old_new"
        left_key, right_key = old_key, new_key
    else:
        left_key = next((k for k in ("product_name", "product name", "name", "product") if k in headers), None)
        right_key = next((k for k in ("sku", "seller_sku", "seller sku", "merchant_sku", "merchant sku") if k in headers), None)
        if not left_key or not right_key:
            raise HTTPException(400, "Use columns Old SKU and New SKU (or legacy Product Name and SKU).")
        mapping_mode = "legacy_product"

    mapping = {}
    for row in rows:
        left = str(row.get(left_key, "") or "").strip()
        right = str(row.get(right_key, "") or "").strip()
        if left and right:
            mapping[left] = right[:100]
    if not mapping:
        raise HTTPException(400, "No usable SKU mapping rows were found in the file.")
    state["mapping"] = mapping
    state["mapping_mode"] = mapping_mode
    state["csv_filename"] = file.filename
    return {"file_id": file_id, "filename": file.filename, "rows": len(mapping), "mapping_mode": mapping_mode}


@app.get("/template-csv")
async def template_csv():
    content = "Old SKU,New SKU\nS-White-Check-SHIrT-(L),W-CH-L\nS- BL-CHk-shirt-M,B-ch-M\nS-check-redshirt-XL,R-ch-XL\nS-shirt BL-CHK-M,B-ch-M\n"
    return Response(content=content, media_type="text/csv", headers={"Content-Disposition": 'attachment; filename="sku_mapping_template.csv"'})


@app.post("/process")
async def process(payload: Dict[str, Any]):
    cleanup_sessions()
    file_id = payload.get("file_id")
    if not file_id or file_id not in sessions:
        raise HTTPException(404, "Upload session not found. Please upload the PDF again.")
    try:
        state = sessions[file_id]
        mapping = state.get("mapping") or {}
        analysis = analyze_pdf(state["original"], mapping, mapping_mode=state.get("mapping_mode", "old_new"))
        first_label = analysis[0]["label_page"] if analysis else 0
        src_doc = fitz.open(stream=state["original"], filetype="pdf")
        first_page = src_doc[first_label]
        page_width = float(first_page.rect.width)
        original_height = float(first_page.rect.height)
        src_doc.close()

        sku_x = float(payload.get("sku_x", page_width * 0.50))
        sku_y = payload.get("sku_y")
        sku_y = float(sku_y) if sku_y is not None else original_height + 32.0
        sku_font_size = float(payload.get("sku_font_size", 10.0))

        output, skus, analysis = process_pdf(
            state["original"],
            mapping=mapping,
            labels_only=True,
            sku_x=sku_x,
            sku_y=sku_y,
            sku_font_size=sku_font_size,
            mapping_mode=state.get("mapping_mode", "old_new"),
        )
        state["processed"] = output
        state["skus"] = skus
        state["analysis"] = analysis
        max_lines = max([len((value or "").splitlines()) for value in skus] or [1])
        sku_area_height = max(60, int(max_lines * (sku_font_size + 3) + 20))
        state["position"] = {"x": sku_x, "y": sku_y, "font_size": sku_font_size, "page_width": page_width, "original_height": original_height, "output_height": original_height + sku_area_height}
        source_doc = fitz.open(stream=state["original"], filetype="pdf")
        source_pages = source_doc.page_count
        source_doc.close()
        return {
            "file_id": file_id,
            "pages": len(skus),
            "source_pages": source_pages,
            "detected": sum(bool(s) for s in skus),
            "missing": sum(not bool(s) for s in skus),
            "skus": skus,
            "analysis": analysis,
            "previews": preview_pdf(output),
            "position": state["position"],
        }
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(500, "Something went wrong while processing the PDF. Please try again.") from exc


@app.post("/generate")
async def generate(payload: GenerateRequest):
    cleanup_sessions()
    state = sessions.get(payload.file_id)
    if not state:
        raise HTTPException(404, "Generated session expired. Please upload the PDF again.")
    analysis = state.get("analysis") or []
    if len(payload.skus) != len(analysis):
        raise HTTPException(400, "SKU count does not match the number of shipping labels.")
    if any(s is not None and any(len(line.strip()) > 100 for line in s.splitlines()) for s in payload.skus):
        raise HTTPException(400, "Each SKU line must be under 100 characters.")
    try:
        position = state.get("position") or {}
        sku_x = float(payload.sku_x)
        sku_y = payload.sku_y if payload.sku_y is not None else position.get("y")
        sku_font_size = float(payload.sku_font_size)
        output, skus, analysis = process_pdf(
            state["original"],
            mapping=state.get("mapping") or {},
            manual_skus=payload.skus,
            labels_only=payload.labels_only,
            sku_x=sku_x,
            sku_y=sku_y,
            sku_font_size=sku_font_size,
            mapping_mode=state.get("mapping_mode", "old_new"),
        )
        state["processed"] = output
        state["skus"] = skus
        state["analysis"] = analysis
        state["position"] = {**position, "x": sku_x, "y": sku_y, "font_size": sku_font_size}
        return {
            "file_id": payload.file_id,
            "pages": len(skus),
            "detected": sum(bool(s) for s in skus),
            "missing": sum(not bool(s) for s in skus),
            "analysis": analysis,
            "previews": preview_pdf(output),
            "position": state["position"],
        }
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(500, "Something went wrong while generating the PDF.") from exc


@app.get("/download/{file_id}")
async def download(file_id: str):
    cleanup_sessions()
    state = sessions.get(file_id)
    if not state or not state.get("processed"):
        raise HTTPException(404, "Generated PDF not found. Please process the PDF again.")
    return Response(
        content=state["processed"],
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="amazon_labels_with_sku.pdf"'},
    )
